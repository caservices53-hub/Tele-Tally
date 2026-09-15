from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any


ALIASES = {
    "date": {"date", "voucher date", "transaction date", "bill date", "invoice date"},
    "voucher_type": {"voucher type", "type", "vch type", "entry type"},
    "party": {"party", "party name", "customer", "supplier", "vendor", "account name", "counterparty"},
    "amount": {"amount", "total", "gross amount", "invoice amount", "transaction amount"},
    "reference": {"reference", "ref", "invoice no", "invoice number", "bill no", "bill number", "document no"},
    "narration": {"narration", "description", "particulars", "remarks", "memo"},
    "debit_ledger": {"debit ledger", "dr ledger", "debit account", "dr account"},
    "credit_ledger": {"credit ledger", "cr ledger", "credit account", "cr account"},
    "bank_ledger": {"bank ledger", "cash bank ledger", "bank account", "bank"},
    "party_ledger": {"party ledger", "tally party ledger"},
    "entries_json": {"entries json", "entries_json", "ledger entries"},
    "items_json": {"items json", "items_json", "inventory entries"},
}


def _norm_header(v: Any) -> str:
    return " ".join(str(v or "").strip().lower().replace("_", " ").split())


def canonicalize(row: dict) -> dict:
    index = {_norm_header(k): v for k, v in row.items()}
    out = {}
    used = set()
    for target, aliases in ALIASES.items():
        for a in aliases:
            if a in index:
                out[target] = index[a]
                used.add(a)
                break
    out["_extra"] = {k: v for k, v in index.items() if k not in used and v not in (None, "")}
    for key in ("party", "reference", "narration", "voucher_type", "debit_ledger", "credit_ledger", "bank_ledger", "party_ledger"):
        if key in out and out[key] is not None:
            out[key] = str(out[key]).strip()
    if "amount" in out and out["amount"] not in (None, ""):
        if isinstance(out["amount"], str):
            out["amount"] = out["amount"].replace(",", "").strip()
        out["amount"] = float(out["amount"])
    if "date" in out and out["date"] is not None:
        d = out["date"]
        if hasattr(d, "strftime"):
            out["date"] = d.strftime("%Y-%m-%d")
        else:
            out["date"] = str(d).strip()
    return out


def load_csv(path: str) -> list[dict]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [canonicalize(r) for r in csv.DictReader(f)]


def load_xlsx(path: str, sheet: str = "") -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("Excel import requires openpyxl. Run: pip install openpyxl") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(x or "") for x in next(rows)]
    except StopIteration:
        return []
    out=[]
    for vals in rows:
        if not any(v not in (None, "") for v in vals):
            continue
        out.append(canonicalize(dict(zip(headers, vals))))
    return out


def load_file(path: str, sheet: str = "") -> list[dict]:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(path)
    if p.suffix.lower() == ".csv":
        return load_csv(str(p))
    if p.suffix.lower() in {".xlsx", ".xlsm"}:
        return load_xlsx(str(p), sheet)
    if p.suffix.lower() == ".json":
        data = json.loads(p.read_text(encoding="utf-8"))
        if not isinstance(data, list):
            raise ValueError("JSON import must contain a list of transaction objects.")
        return [canonicalize(x) for x in data]
    raise ValueError("Supported bulk files: .csv, .xlsx, .xlsm, .json")
